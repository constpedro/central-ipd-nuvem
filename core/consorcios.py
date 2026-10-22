# -*- coding: utf-8 -*-
"""
Relatório Geral de Produção PPI por consórcio.

Reproduz o relatório que era montado à mão (Relatorio_<CONSORCIO>_Final.xlsx):
para cada município da lista, consulta os subgrupos 0203 e 0202 do PPI e
classifica:

    Citopatológicos    = Forma de Organização 020301 (0203010xxx)
    Anatopatológicos   = Forma de Organização 020302 (0203020xxx)
    Análises Clínicas  = subgrupo 0202 inteiro (só valor; o PPI não discrimina
                         quantidade útil por procedimento nesse subgrupo)
    Atend. <polo>      = parcela do valor anual cujo município de atendimento
                         é o polo (Varginha por padrão)
    Mensal             = anual ÷ 12

Sai uma tabela na tela e um .xlsx com três abas (Resumo, Detalhe <polo>, Notas),
no mesmo layout dos relatórios anteriores.
"""
import time
import copy
import json
import re
import uuid
from datetime import datetime

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from core import ppi_engine

PASTA = config.DATA_DIR / "consorcios"


def categoria(subgrupo, codigo):
    if subgrupo == "0202":
        return "Análises Clínicas"
    return "Citopatológico" if codigo.startswith(config.PREFIXO_CITO) else "Anatopatológico"


def codigo_atendimento(rotulo):
    return rotulo.split(" - ", 1)[0].strip()


def agregar_cidade(dados, polo_codigo):
    """Uma linha do resumo + linhas de detalhe do polo, a partir do cache de uma cidade."""
    linha = {
        "municipio": dados["cidade"],
        "codigo_ibge": dados["codigo_ibge"],
        "anato_qtd": 0.0, "anato_valor": 0.0,
        "cito_qtd": 0.0, "cito_valor": 0.0,
        "ac_valor": 0.0,
        "polo_anato": 0.0, "polo_cito": 0.0, "polo_ac": 0.0,
        "conferencia": [],
    }
    detalhe = []
    for sub, bloco in dados["subgrupos"].items():
        soma_qtd = soma_val = 0.0
        for p in bloco["procedimentos"]:
            cat = categoria(sub, p["codigo"])
            soma_qtd += p["qtd_anual"]
            soma_val += p["valor_anual"]
            no_polo = codigo_atendimento(p["municipio_atendimento"]) == polo_codigo
            if cat == "Citopatológico":
                linha["cito_qtd"] += p["qtd_anual"]
                linha["cito_valor"] += p["valor_anual"]
                if no_polo:
                    linha["polo_cito"] += p["valor_anual"]
            elif cat == "Anatopatológico":
                linha["anato_qtd"] += p["qtd_anual"]
                linha["anato_valor"] += p["valor_anual"]
                if no_polo:
                    linha["polo_anato"] += p["valor_anual"]
            else:
                linha["ac_valor"] += p["valor_anual"]
                if no_polo:
                    linha["polo_ac"] += p["valor_anual"]
            if no_polo:
                detalhe.append({
                    "municipio": dados["cidade"], "categoria": cat, "codigo": p["codigo"],
                    "procedimento": p["nome"], "qtd": p["qtd_anual"], "valor": p["valor_anual"],
                })
        oficial = bloco.get("total_oficial")
        if oficial:
            bate = abs(soma_val - oficial[1]) < 0.01
            linha["conferencia"].append({"subgrupo": sub, "soma": round(soma_val, 2), "oficial": oficial[1], "ok": bate})
    linha["polo_total"] = linha["polo_anato"] + linha["polo_cito"] + linha["polo_ac"]
    linha["total_anual"] = linha["anato_valor"] + linha["cito_valor"] + linha["ac_valor"]
    linha["total_mensal"] = linha["total_anual"] / 12
    return linha, detalhe


def executar(job):
    p = job.params
    nome, cidades, competencia = p["nome"], p["cidades"], p["competencia"]
    polo_codigo, polo_nome = p["polo_codigo"], p["polo_nome"]
    session = ppi_engine.nova_sessao()
    job.progresso(0, len(cidades))
    job.definir_itens(cidades)
    linhas, detalhe, falhas = [], [], []
    for i, cidade in enumerate(cidades, 1):
        if job.cancelado:
            job.escrever("Cancelado — nenhum relatório foi salvo (as consultas feitas ficam no cache).")
            return {"arquivos": []}
        job.marcar(cidade, "rodando")
        comecou = time.monotonic()
        try:
            dados = ppi_engine.consultar_cidade(session, cidade, competencia, log=job.escrever)
            linha, det = agregar_cidade(dados, polo_codigo)
            linhas.append(linha)
            detalhe.extend(det)
            job.marcar(cidade, "ok", time.monotonic() - comecou)
        except Exception as e:  # noqa: BLE001 — uma cidade com erro não derruba o consórcio
            job.escrever(f"[ERRO] {cidade}: {e}")
            falhas.append({"municipio": cidade, "erro": str(e)})
            job.marcar(cidade, "erro", time.monotonic() - comecou)
        job.progresso(i, len(cidades))

    totais = {k: sum(l[k] for l in linhas) for k in (
        "anato_qtd", "anato_valor", "cito_qtd", "cito_valor", "ac_valor",
        "polo_anato", "polo_cito", "polo_ac", "polo_total", "total_anual", "total_mensal")}
    divergencias = [
        {"municipio": l["municipio"], **c} for l in linhas for c in l["conferencia"] if not c["ok"]
    ]
    relatorio = {
        "id": uuid.uuid4().hex[:8],
        "nome": nome,
        "competencia": competencia,
        "competencia_fmt": ppi_engine.formatar_competencia(competencia),
        "polo_codigo": polo_codigo,
        "polo_nome": polo_nome,
        "gerado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "n_municipios": len(cidades),
        "linhas": linhas,
        "detalhe": detalhe,
        "falhas": falhas,
        "totais": totais,
        "divergencias": divergencias,
    }
    PASTA.mkdir(parents=True, exist_ok=True)
    slug = re.sub(r"[^a-z0-9]+", "_", nome.lower()).strip("_") or "consorcio"
    base = f"Relatorio_{slug.upper()}_{competencia}_{relatorio['id']}"
    relatorio["arquivo_xlsx"] = base + ".xlsx"
    (PASTA / (base + ".json")).write_text(json.dumps(relatorio, ensure_ascii=False, indent=1), encoding="utf-8")
    exportar_xlsx(relatorio, PASTA / relatorio["arquivo_xlsx"])
    job.escrever(f"Planilha salva: {relatorio['arquivo_xlsx']}")
    return {"relatorio_id": relatorio["id"], "arquivo_xlsx": relatorio["arquivo_xlsx"], "n_falhas": len(falhas)}


_resumos = {}  # nome do .json -> (mtime, resumo): cada relatório é lido uma vez só


def contar():
    return sum(1 for _ in PASTA.glob("*.json")) if PASTA.is_dir() else 0


def listar():
    if not PASTA.is_dir():
        return []
    arquivos = sorted(((p, p.stat().st_mtime) for p in PASTA.glob("*.json")), key=lambda x: x[1], reverse=True)
    saida = []
    for p, mtime in arquivos:
        guardado = _resumos.get(p.name)
        if guardado and guardado[0] == mtime:
            saida.append(guardado[1])
            continue
        try:
            r = json.loads(p.read_text(encoding="utf-8"))
            resumo = ({k: r[k] for k in ("id", "nome", "competencia_fmt", "polo_nome", "gerado_em", "n_municipios", "arquivo_xlsx")}
                      | {"total_anual": r["totais"]["total_anual"], "polo_total": r["totais"]["polo_total"]})
        except (OSError, ValueError, KeyError, TypeError):
            continue  # arquivo incompleto ou de formato antigo: não derruba a página
        _resumos[p.name] = (mtime, resumo)
        saida.append(resumo)
    return saida


def carregar(relatorio_id):
    if not re.fullmatch(r"[0-9a-f]{8}", relatorio_id or ""):
        return None
    for p in PASTA.glob(f"*_{relatorio_id}.json"):
        return json.loads(p.read_text(encoding="utf-8"))
    return None


# --- exportação .xlsx ------------------------------------------------------------

FINO = Side(style="thin", color="D5D1CB")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
CABECALHO = PatternFill("solid", fgColor="1D1B19")
CABECALHO_2 = PatternFill("solid", fgColor="F0EEE6")
DESTAQUE = PatternFill("solid", fgColor="FBE9EA")
BRANCO = Font(bold=True, color="FFFFFF")
MOEDA = '"R$" #,##0.00'
QTD = "#,##0.0"


def _modelo_proprio(nome):
    """Modelo visual oficial do próprio consórcio, ou None se ele não tiver um."""
    compactado = re.sub(r"[^a-z0-9]+", "", str(nome).lower())
    for chave, caminho in getattr(config, "RELATORIO_MODELOS", {}).items():
        if re.sub(r"[^a-z0-9]+", "", chave.lower()) in compactado and caminho.is_file():
            return caminho
    return None


def _modelo_relatorio(nome):
    """Retorna o modelo visual oficial correspondente ao consórcio."""
    proprio = _modelo_proprio(nome)
    if proprio:
        return proprio
    padrao = getattr(config, "RELATORIO_MODELO_PADRAO", None)
    return padrao if padrao and padrao.is_file() else None


def _achar_total(ws, max_col=None):
    max_col = max_col or ws.max_column
    for row in range(1, ws.max_row + 1):
        for col in range(1, max_col + 1):
            valor = ws.cell(row=row, column=col).value
            if isinstance(valor, str) and valor.strip().upper().startswith("TOTAL"):
                return row
    return None


def _copiar_estilo_linha(ws, origem, destino, max_col):
    """Copia somente aparência/altura para linhas novas de um modelo."""
    if origem == destino:
        return
    ws.row_dimensions[destino].height = ws.row_dimensions[origem].height
    ws.row_dimensions[destino].hidden = ws.row_dimensions[origem].hidden
    for col in range(1, max_col + 1):
        fonte = ws.cell(row=origem, column=col)
        alvo = ws.cell(row=destino, column=col)
        if fonte.has_style:
            alvo._style = copy.copy(fonte._style)
        if fonte.number_format:
            alvo.number_format = fonte.number_format
        alvo.alignment = copy.copy(fonte.alignment)
        alvo.font = copy.copy(fonte.font)
        alvo.fill = copy.copy(fonte.fill)
        alvo.border = copy.copy(fonte.border)
        alvo.protection = copy.copy(fonte.protection)


def _redimensionar_antes_do_total(ws, total_atual, total_desejado, linha_modelo, max_col):
    """Insere/remove linhas de detalhe e remonta mesclagens abaixo do total."""
    if not total_atual or total_atual == total_desejado:
        return

    mesclagens = []
    for faixa in list(ws.merged_cells.ranges):
        if faixa.min_row >= total_atual:
            mesclagens.append((faixa.min_row, faixa.min_col, faixa.max_row, faixa.max_col))
            ws.unmerge_cells(str(faixa))

    delta = total_desejado - total_atual
    if delta > 0:
        ws.insert_rows(total_atual, delta)
        for row in range(total_atual, total_desejado):
            _copiar_estilo_linha(ws, linha_modelo, row, max_col)
    else:
        ws.delete_rows(total_desejado, -delta)

    for min_row, min_col, max_row, max_col_faixa in mesclagens:
        ws.merge_cells(
            start_row=min_row + delta,
            start_column=min_col,
            end_row=max_row + delta,
            end_column=max_col_faixa,
        )


def _valor(v, casas=2):
    return round(float(v or 0), casas)


def _destacar_polo(celula, vizinha, valor):
    """O destaque do polo vem pintado no modelo por posição; refaz conforme o valor real."""
    if valor:
        celula.fill = PatternFill("solid", fgColor="FFFFD9D9")
        celula.font = Font(name=vizinha.font.name, size=vizinha.font.sz, bold=True, color="FF8B0000")
    else:
        celula.fill = copy.copy(vizinha.fill)
        celula.font = copy.copy(vizinha.font)


def _nota_resumo(r):
    polo = r["polo_nome"]
    return (
        "Citopatológicos = PPI 020301  |  Anatopatológicos = PPI 020302  |  "
        "Análises Clínicas = PPI 0202 (valor total)  |  "
        "Qtd./Valor Mensal = anual ÷ 12 (média programada)  |  "
        f"Atend. {polo} = soma da parcela de cada categoria atendida em {polo} "
        f'(ver aba "Detalhe {polo}")  |  Fonte: PPI Assistencial SES/MG, '
        f"competência {r['competencia_fmt']}."
    )


def _exportar_xlsx_modelo(r, caminho, modelo):
    """Preenche o modelo oficial sem recriar o layout do relatório."""
    wb = load_workbook(modelo)
    ws = wb.worksheets[0]
    polo = r["polo_nome"]
    t = r["totais"]
    linhas = r["linhas"]

    # Mantém o nome original nos quatro modelos fornecidos; nos demais, a aba
    # recebe o nome do consórcio. No fallback (consórcio sem modelo, que usa o
    # do CISMIP) a aba não pode continuar "Relatorio CISMIP".
    titulo_aba = str(ws.title).lower()
    original = titulo_aba.startswith("relatorio ") and any(
        c in titulo_aba for c in ("cismip", "cisamesp", "cissul", "cismas"))
    if not original or _modelo_proprio(r["nome"]) != modelo:
        ws.title = re.sub(r"[\\/*?:\[\]]", "-", f"Relatorio {r['nome']}")[:31]

    ws["A2"] = f"Relatório Geral de Produção PPI — Consórcio {r['nome']}"
    ws["A3"] = (
        f"Competência: {r['competencia_fmt']}  |  {len(linhas)} municípios consorciados — "
        "valores e quantidades anuais e mensais (méd.) programados no PPI, "
        f"com destaque para atendimentos em {polo}."
    )
    ws["A5"] = f"VALOR TOTAL ATENDIDO EM {polo.upper()}"
    for celula, rotulo in (("B6", "Anatopatológicos"), ("E6", "Citopatológicos"),
                           ("H6", "Análises Clínicas"), ("K6", f"TOTAL GERAL {polo.upper()}")):
        ws[celula] = rotulo
    for celula, valor in (("B7", t["polo_anato"]), ("E7", t["polo_cito"]),
                          ("H7", t["polo_ac"]), ("K7", t["polo_total"])):
        ws[celula] = _valor(valor)
    ws["A8"] = f'Ver detalhamento completo por cidade e procedimento na aba "Detalhe {polo}".'

    # O padrão oficial usa duas linhas de cabeçalho e inicia os municípios na
    # linha 12. As mesclagens e estilos continuam vindo do arquivo-modelo.
    cabecalhos = {
        "A10": "Município",
        "B10": "Anatopatológicos",
        "F10": "Citopatológicos",
        "J10": "Análises Clínicas",
        "L10": f"Atend. {polo} (R$)",
        "M10": "Total Geral Anual (R$)",
        "N10": "Total Geral Mensal (R$)",
        "B11": "Qtd. Anual", "C11": "Valor Anual", "D11": "Qtd. Mensal", "E11": "Valor Mensal",
        "F11": "Qtd. Anual", "G11": "Valor Anual", "H11": "Qtd. Mensal", "I11": "Valor Mensal",
        "J11": "Valor Anual", "K11": "Valor Mensal",
    }
    for celula, valor in cabecalhos.items():
        ws[celula] = valor

    linha_inicio = 12
    # No resumo, os cabeçalhos M10/N10 contêm "Total Geral". O total real é
    # identificado exclusivamente na coluna A.
    linha_total_atual = _achar_total(ws, 1)
    linha_total = linha_inicio + len(linhas)
    _redimensionar_antes_do_total(ws, linha_total_atual, linha_total, linha_inicio, 14)

    for i, linha in enumerate(linhas, start=linha_inicio):
        valores = {
            1: linha["municipio"],
            2: _valor(linha["anato_qtd"], 4),
            3: _valor(linha["anato_valor"]),
            4: f"=B{i}/12",
            5: f"=C{i}/12",
            6: _valor(linha["cito_qtd"], 4),
            7: _valor(linha["cito_valor"]),
            8: f"=F{i}/12",
            9: f"=G{i}/12",
            10: _valor(linha["ac_valor"]),
            11: f"=J{i}/12",
            12: _valor(linha["polo_total"]),
            13: f"=C{i}+G{i}+J{i}",
            14: f"=M{i}/12",
        }
        for col, valor in valores.items():
            ws.cell(row=i, column=col).value = valor
        _destacar_polo(ws.cell(row=i, column=12), ws.cell(row=i, column=13), linha["polo_total"])

    rotulo_total_modelo = ws.cell(row=linha_total, column=1).value
    if isinstance(rotulo_total_modelo, str) and rotulo_total_modelo.strip().upper() == "TOTAL":
        ws.cell(row=linha_total, column=1).value = "TOTAL"
    else:
        ws.cell(row=linha_total, column=1).value = f"TOTAL {r['nome']} ({len(linhas)} municípios)"
    for col in range(2, 15):
        letra = get_column_letter(col)
        ws.cell(row=linha_total, column=col).value = (
            f"=SUM({letra}{linha_inicio}:{letra}{linha_total - 1})" if linhas else 0
        )

    # Alguns modelos têm uma nota curta no rodapé do resumo. Atualiza apenas
    # o texto existente, preservando sua mesclagem e formatação.
    for row in range(linha_total + 1, ws.max_row + 1):
        valor = ws.cell(row=row, column=1).value
        if isinstance(valor, str) and valor.strip().lower().startswith("anatopat"):
            ws.cell(row=row, column=1).value = _nota_resumo(r)

    # --- Detalhe do polo: a posição do cabeçalho varia entre os modelos.
    wd = wb.worksheets[1]
    wd.title = f"Detalhe {polo}"[:31]
    linha_cab = None
    for row in range(1, wd.max_row + 1):
        if str(wd.cell(row=row, column=1).value or "").strip().lower() == "município origem":
            linha_cab = row
            break
    if linha_cab is None:
        linha_cab = 3
    linha_det_inicio = linha_cab + 1
    linha_det_total_atual = _achar_total(wd, 6)
    coluna_total_detalhe = 4
    if linha_det_total_atual:
        for col in range(1, 7):
            if str(wd.cell(row=linha_det_total_atual, column=col).value or "").strip().upper() == "TOTAL":
                coluna_total_detalhe = col
                break
    linha_det_total = linha_det_inicio + len(r["detalhe"])

    # Heurísticas de formatação lidas das linhas-modelo do template. Precisam
    # vir ANTES do redimensionamento: depois dele as linhas mudam de lugar (ou
    # somem, quando o detalhe está vazio) e os índices apontariam para o lugar errado.
    modelo_codigo = str(wd.cell(row=linha_det_inicio, column=3).value or "")
    codigo_longo = len(modelo_codigo.replace(".", "")) >= 8
    usa_traco_qtd = any(
        str(wd.cell(row=row, column=2).value or "").upper().startswith("ANÁLISES")
        and wd.cell(row=row, column=5).value == "-"
        for row in range(linha_det_inicio, min(linha_det_total_atual or wd.max_row, wd.max_row) + 1)
    )

    if linha_det_total_atual:
        _redimensionar_antes_do_total(wd, linha_det_total_atual, linha_det_total, linha_det_inicio, 6)
    if linha_cab > 1:
        wd["A1"] = f"Detalhamento — Procedimentos atendidos em {polo}"
    cab_detalhe = ["Município Origem", "Categoria", "Código", "Procedimento", "Qtd. Anual", "Valor Anual (R$)"]
    for col, valor in enumerate(cab_detalhe, start=1):
        wd.cell(row=linha_cab, column=col).value = valor

    detalhes = sorted(r["detalhe"], key=lambda x: (x["municipio"], x["categoria"], x["codigo"]))
    for row, item in enumerate(detalhes, start=linha_det_inicio):
        codigo = str(item.get("codigo") or "")
        if not codigo_longo and len(codigo) > 4:
            codigo = "0202" if item["categoria"] == "Análises Clínicas" else codigo[-3:]
        qtd = item.get("qtd")
        if qtd is None and usa_traco_qtd:
            qtd = "-"
        wd.cell(row=row, column=1).value = item["municipio"]
        wd.cell(row=row, column=2).value = item["categoria"].upper()
        wd.cell(row=row, column=3).value = codigo
        wd.cell(row=row, column=4).value = item["procedimento"]
        wd.cell(row=row, column=5).value = None if qtd is None else (_valor(qtd, 4) if isinstance(qtd, (int, float)) else qtd)
        wd.cell(row=row, column=6).value = _valor(item["valor"])
    for col in range(1, 6):
        cell = wd.cell(row=linha_det_total, column=col)
        if not isinstance(cell, MergedCell):
            cell.value = "TOTAL" if col == coluna_total_detalhe else None
    wd.cell(row=linha_det_total, column=6).value = (
        f"=SUM(F{linha_det_inicio}:F{linha_det_total - 1})" if detalhes else 0
    )

    # --- Notas: substitui apenas o conteúdo, mantendo a aba e sua aparência.
    wn = wb.worksheets[2]
    notas = [
        "Notas e premissas",
        f"Consórcio {r['nome']} — competência {r['competencia_fmt']} — {len(linhas)} municípios. Gerado pela Central IPD em {r['gerado_em']}.",
        "Fonte dos dados: PPI Assistencial / SES-MG (ppiassistencial.saude.mg.gov.br), subgrupo 0203 e subgrupo 0202.",
        "Citopatológicos = Forma de Organização 020301 (códigos 0203010xxx). Anatopatológicos = Forma de Organização 020302 (códigos 0203020xxx).",
        "Análises Clínicas = valor anual total do PPI subgrupo 0202 — sem discriminação de quantidade.",
        "Qtd./Valor Mensal = anual ÷ 12 (média mensal programada; o PPI não fornece quebra mensal real).",
        f"Atend. {polo} = parcela do valor anual de cada categoria cujo Município Atendimento é {polo} ({r['polo_codigo']}).",
        f"Polo de destaque: {polo} — padrão adotado independentemente da sede formal do consórcio.",
        "Municípios: " + ", ".join(l["municipio"] for l in linhas) + ".",
    ]
    if r["divergencias"]:
        notas.append("ATENÇÃO — divergências entre a soma das linhas extraídas e o total oficial: " + "; ".join(
            f"{d['municipio']} ({d['subgrupo']}: soma R$ {d['soma']:.2f} × oficial R$ {d['oficial']:.2f})" for d in r["divergencias"]))
    else:
        notas.append("Conferência: a soma das linhas extraídas foi comparada com o total oficial de cada município e subgrupo — 0 divergências.")
    if r["falhas"]:
        notas.append("Municípios NÃO incluídos por erro na consulta: " + "; ".join(
            f"{f['municipio']} ({f['erro']})" for f in r["falhas"]))
    for row in range(1, wn.max_row + 1):
        wn.cell(row=row, column=1).value = None
    for row, nota in enumerate(notas, start=1):
        if row > wn.max_row:
            _copiar_estilo_linha(wn, wn.max_row, row, 1)
        wn.cell(row=row, column=1).value = nota

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except AttributeError:
        pass
    wb.save(caminho)


def exportar_xlsx(r, caminho):
    modelo = _modelo_relatorio(r["nome"])
    if modelo:
        _exportar_xlsx_modelo(r, caminho, modelo)
        return
    _exportar_xlsx_generico(r, caminho)


def _exportar_xlsx_generico(r, caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    polo = r["polo_nome"]
    t = r["totais"]

    ws["A1"] = f"Relatório Geral de Produção PPI — Consórcio {r['nome']}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Competência {r['competencia_fmt']} — {r['n_municipios']} municípios — valores e quantidades "
                f"anuais e mensais (média) programados no PPI, com destaque para atendimentos em {polo}.")
    ws["A4"] = f"VALOR TOTAL ATENDIDO EM {polo.upper()}"
    ws["A4"].font = Font(bold=True)
    for col, (rot, val) in enumerate([("Anatopatológicos", t["polo_anato"]), ("Citopatológicos", t["polo_cito"]),
                                      ("Análises Clínicas", t["polo_ac"]), (f"TOTAL GERAL {polo.upper()}", t["polo_total"])], start=1):
        c = ws.cell(row=5, column=col, value=rot)
        c.font = Font(bold=True)
        c.fill = DESTAQUE
        c.border = BORDA
        v = ws.cell(row=6, column=col, value=round(val, 2))
        v.number_format = MOEDA
        v.border = BORDA
    ws["A7"] = f'Ver detalhamento por cidade e procedimento na aba "Detalhe {polo}".'
    ws["A7"].font = Font(italic=True, color="7B766F")

    lin = 9
    grupos = [("Município", 1), ("Anatopatológicos", 4), ("Citopatológicos", 4), ("Análises Clínicas", 2),
              (f"Atend. {polo} (R$)", 1), ("Total Geral Anual (R$)", 1), ("Total Geral Mensal (R$)", 1)]
    col = 1
    for rot, largura in grupos:
        ws.cell(row=lin, column=col, value=rot)
        if largura > 1:
            ws.merge_cells(start_row=lin, start_column=col, end_row=lin, end_column=col + largura - 1)
        col += largura
    sub = ["", "Qtd. Anual", "Valor Anual", "Qtd. Mensal", "Valor Mensal",
           "Qtd. Anual", "Valor Anual", "Qtd. Mensal", "Valor Mensal", "Valor Anual", "Valor Mensal", "", "", ""]
    for i, s in enumerate(sub, start=1):
        ws.cell(row=lin + 1, column=i, value=s or None)
    for rr in (lin, lin + 1):
        for i in range(1, 15):
            c = ws.cell(row=rr, column=i)
            c.font = BRANCO if rr == lin else Font(bold=True)
            c.fill = CABECALHO if rr == lin else CABECALHO_2
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDA

    formatos = [None, QTD, MOEDA, QTD, MOEDA, QTD, MOEDA, QTD, MOEDA, MOEDA, MOEDA, MOEDA, MOEDA, MOEDA]
    lin += 2
    for l in r["linhas"]:
        vals = [l["municipio"], l["anato_qtd"], l["anato_valor"], l["anato_qtd"] / 12, l["anato_valor"] / 12,
                l["cito_qtd"], l["cito_valor"], l["cito_qtd"] / 12, l["cito_valor"] / 12,
                l["ac_valor"], l["ac_valor"] / 12,
                l["polo_total"] if l["polo_total"] else None, l["total_anual"], l["total_mensal"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=lin, column=i, value=(round(v, 4) if isinstance(v, float) else v))
            if formatos[i - 1]:
                c.number_format = formatos[i - 1]
            if i == 12 and v is None:
                c.value = "-"
                c.alignment = Alignment(horizontal="center")
            c.border = BORDA
        lin += 1
    vals = ["TOTAL", t["anato_qtd"], t["anato_valor"], t["anato_qtd"] / 12, t["anato_valor"] / 12,
            t["cito_qtd"], t["cito_valor"], t["cito_qtd"] / 12, t["cito_valor"] / 12,
            t["ac_valor"], t["ac_valor"] / 12, t["polo_total"], t["total_anual"], t["total_mensal"]]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=lin, column=i, value=(round(v, 4) if isinstance(v, float) else v))
        c.font = Font(bold=True)
        c.fill = CABECALHO_2
        c.border = BORDA
        if formatos[i - 1]:
            c.number_format = formatos[i - 1]
    larguras = [30, 11, 14, 11, 14, 11, 14, 11, 14, 15, 14, 18, 18, 18]
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=11, column=2)

    # --- Detalhe do polo
    wd = wb.create_sheet(f"Detalhe {polo}"[:31])
    wd["A1"] = f"Detalhamento — Atendimento em {polo} ({r['polo_codigo']})"
    wd["A1"].font = Font(bold=True, size=13)
    wd["A2"] = f"{r['nome']} — Competência {r['competencia_fmt']}"
    cab = ["Município Origem", "Categoria", "Código", "Procedimento", "Qtd. Anual", "Valor Anual (R$)"]
    for i, s in enumerate(cab, start=1):
        c = wd.cell(row=4, column=i, value=s)
        c.font = BRANCO
        c.fill = CABECALHO
        c.border = BORDA
    lin = 5
    for d in sorted(r["detalhe"], key=lambda x: (x["municipio"], x["categoria"], x["codigo"])):
        vals = [d["municipio"], d["categoria"], d["codigo"], d["procedimento"], d["qtd"], d["valor"]]
        for i, v in enumerate(vals, start=1):
            c = wd.cell(row=lin, column=i, value=v)
            c.border = BORDA
            if i == 5:
                c.number_format = QTD
            if i == 6:
                c.number_format = MOEDA
        lin += 1
    for i, w in enumerate([28, 18, 14, 60, 12, 16], start=1):
        wd.column_dimensions[get_column_letter(i)].width = w
    wd.freeze_panes = "A5"

    # --- Notas
    wn = wb.create_sheet("Notas")
    notas = [
        "Notas e premissas",
        f"Consórcio {r['nome']} — competência {r['competencia_fmt']} — {r['n_municipios']} municípios. Gerado pela Central IPD em {r['gerado_em']}.",
        "Fonte dos dados: PPI Assistencial / SES-MG (ppiassistencial.saude.mg.gov.br), subgrupo 0203 (Diagnóstico por anatomia patológica e citopatologia) e subgrupo 0202 (Diagnóstico em laboratório clínico).",
        "Citopatológicos = Forma de Organização 020301 (códigos 0203010xxx). Anatopatológicos = Forma de Organização 020302 (códigos 0203020xxx).",
        "Análises Clínicas = valor anual total do PPI subgrupo 0202 — sem discriminação de quantidade (o PPI não traz quantidade útil por procedimento nesse subgrupo).",
        "Qtd./Valor Mensal = anual ÷ 12 (média mensal programada; o PPI não fornece quebra mensal real).",
        f"Atend. {polo} = parcela do valor anual de cada categoria cujo Município Atendimento é {polo} ({r['polo_codigo']}). Detalhamento na aba \"Detalhe {polo}\".",
        f"Polo de destaque: {polo} — padrão adotado independentemente da sede formal do consórcio.",
        "Municípios: " + ", ".join(l["municipio"] for l in r["linhas"]) + ".",
    ]
    if r["divergencias"]:
        notas.append("ATENÇÃO — divergências entre a soma das linhas extraídas e o total oficial do site: " + "; ".join(
            f"{d['municipio']} ({d['subgrupo']}: soma R$ {d['soma']:.2f} × oficial R$ {d['oficial']:.2f})" for d in r["divergencias"]))
    else:
        notas.append("Conferência: para cada município e subgrupo, a soma das linhas extraídas foi comparada com a linha Total oficial do próprio relatório do PPI — 0 divergências.")
    if r["falhas"]:
        notas.append("Municípios NÃO incluídos por erro na consulta: " + "; ".join(f"{f['municipio']} ({f['erro']})" for f in r["falhas"]))
    for i, n in enumerate(notas, start=1):
        c = wn.cell(row=i, column=1, value=n)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)
    wn.column_dimensions["A"].width = 140
    wb.save(caminho)

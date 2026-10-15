# -*- coding: utf-8 -*-
"""
Teto PPI de Varginha: monta a planilha mensal "TETO PPI - <MÊS>.xlsx".

Todo mês a planilha era montada à mão no PPI Assistencial com:
  Foco de Pesquisa = Atendimento, Atendimento = 317070 - Varginha,
  Discriminar por = Gestão/Município Origem, Grupo 02,
e uma consulta por subgrupo (0203 e 0202) e por procedimento do 0203.

Aqui fazemos as mesmas consultas e regravamos só as abas de dados de uma
planilha-modelo (a última usada). As abas de fórmulas (SUS AP, SUS COLPO,
Geral0203) são refeitas com o número de linhas do mês e as abas ShiftAP /
ShiftColpo — com os limites digitados no Shift — ficam como estão no modelo.
"""
import re
import shutil
import time
from datetime import datetime

import openpyxl
import requests
from openpyxl.worksheet.views import Selection
from bs4 import BeautifulSoup

import config
from core import ppi_engine, xlsx_calc

PASTA = config.DATA_DIR / "teto_ppi"
MODELO = PASTA / "modelo.xlsx"
SAIDA = PASTA / "gerados"

ATENDIMENTO = "Varginha"
# (forma de organização, código do procedimento) — uma aba por procedimento,
# na mesma ordem das colunas C:I da aba Geral0203.
PROCEDIMENTOS = [
    ("020301", "0203010019"),
    ("020301", "0203010043"),
    ("020301", "0203010086"),
    ("020302", "0203020022"),
    ("020302", "0203020065"),
    ("020302", "0203020073"),
    ("020302", "0203020081"),
]
MESES = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
FONTE = "Fonte: SES/ MG - Diretoria de Programação Pactuada Integrada (DPPI)."
_MUNICIPIO = re.compile(r"^\d{6}\s*-\s*\S")


def _num(txt):
    txt = txt.replace("R$", "").strip().replace(".", "").replace(",", ".")
    try:
        return float(txt) if txt else 0.0
    except ValueError:
        return 0.0


def parse_por_origem(html):
    """Linhas 'Município de Origem | Qtd | Valor | SADT' -> [(municipio, qtd, valor, sadt)].

    Se a página trouxer mais de um bloco (gestão municipal e estadual), o mesmo
    município aparece em mais de um: somamos, como faria a soma na planilha.
    """
    somas = {}
    for tr in BeautifulSoup(html, "html.parser").find_all("tr"):
        tds = [td.get_text(" ", strip=True) for td in tr.find_all("td")]
        if len(tds) < 4 or not _MUNICIPIO.match(tds[0]):
            continue
        nome = re.sub(r"\s*-\s*", " - ", tds[0], count=1)
        q, v, s = somas.get(nome, (0.0, 0.0, 0.0))
        somas[nome] = (q + _num(tds[1]), v + _num(tds[2]), s + _num(tds[3]))
    linhas = [(m, *vals) for m, vals in somas.items()]
    linhas.sort(key=lambda l: l[0])
    return linhas


def total_oficial(html):
    """Totais impressos pelo site, separados por tipo de bloco.

    A página pode ter blocos de "Município de Origem", "Microrregião de Origem"
    e "Estado de Origem" (cada um com seu Total Recurso), por complexidade.
    A planilha usa só os municípios, então devolvemos:
      {"municipio": (qtd, valor) ou None, "outros": [(rotulo, qtd, valor), ...], "geral": (qtd, valor) ou None}
    """
    bloco, mun, achou, outros, geral = None, [0.0, 0.0], False, [], None
    for tr in BeautifulSoup(html, "html.parser").find_all("tr"):
        tds = [td.get_text(" ", strip=True) for td in tr.find_all(["td", "th"])]
        if not tds:
            continue
        rot = tds[0].lower()
        if rot.endswith("de origem"):
            bloco = tds[0]
        elif rot.startswith("total recurso") and len(tds) >= 3:
            if bloco and bloco.lower().startswith("munic"):
                mun[0] += _num(tds[1]); mun[1] += _num(tds[2]); achou = True
            elif bloco:
                outros.append((bloco, _num(tds[1]), _num(tds[2])))
            bloco = None
        elif rot.startswith("total geral") and len(tds) >= 3:
            geral = (_num(tds[1]), _num(tds[2]))
    return {"municipio": tuple(mun) if achou else None, "outros": outros, "geral": geral}


def conferir(nome, linhas, totais):
    """Compara a soma das linhas extraídas com o total do site; erro se divergir."""
    total = totais["municipio"]
    if total is None and not linhas:
        return 0.0, 0.0
    if total is None:
        raise RuntimeError(f"{nome}: o site não trouxe a linha Total Recurso — não dá para conferir")
    qtd = sum(l[1] for l in linhas)
    val = sum(l[2] for l in linhas)
    if abs(qtd - total[0]) > 0.5 or abs(val - total[1]) > 0.05:
        raise RuntimeError(
            f"{nome}: soma das linhas (qtd {qtd:,.0f} · R$ {val:,.2f}) não bate com o total do site "
            f"(qtd {total[0]:,.0f} · R$ {total[1]:,.2f}) — planilha não gerada"
            .replace(",", "X").replace(".", ",").replace("X", "."))
    return qtd, val


def competencias_publicadas(sessao=None, base=None):
    """Competências que o PPI oferece no formulário, da mais recente para a mais antiga."""
    try:
        sessao = sessao or ppi_engine.nova_sessao()
        base = base or ppi_engine._motor().BASE + "/"
        html = sessao.get(base + "consulta-ppi", timeout=20).text
        sel = BeautifulSoup(html, "html.parser").find("select", {"name": "competencia"})
        return [o.get("value") for o in sel.find_all("option") if (o.get("value") or "").isdigit()] if sel else []
    except Exception:  # noqa: BLE001 — sem rede, segue sem a checagem
        return []


class Consulta:
    """Sessão no PPI com o Atendimento já resolvido."""

    def __init__(self, competencia):
        self.g = ppi_engine._motor()
        self.base = self.g.BASE + "/"
        self.competencia = competencia
        self.s = ppi_engine.nova_sessao()
        self.g.garantir_sessao(self.s)
        disponiveis = competencias_publicadas(self.s, self.base)
        if disponiveis and competencia not in disponiveis:
            ultima = disponiveis[0]
            raise RuntimeError(
                f"a competência {ppi_engine.formatar_competencia(competencia)} ainda não foi publicada no PPI — "
                f"a mais recente disponível é {ppi_engine.formatar_competencia(ultima)} ({ultima})")
        r = self.s.post(self.base + "consulta-ppi/autocomplete",
                        data={"dataType": "json", "atendimento": ATENDIMENTO}, timeout=30)
        r.raise_for_status()
        opcoes = [o for o in r.json() if o["name"].split(" - ", 1)[-1].strip().lower() == ATENDIMENTO.lower()]
        if not opcoes:
            raise RuntimeError(f"o PPI não encontrou o atendimento {ATENDIMENTO}")
        self.atendimento, self.atendimento_id = opcoes[0]["name"], opcoes[0]["code"]
        self._ids = {}

    def id_procedimento(self, forma, codigo):
        """O formulário manda o id interno do procedimento, não o código SIGTAP."""
        if forma not in self._ids:
            r = self.s.get(self.base + "consulta-ppi/ajax-busca-procedimentos",
                           params={"forma_organizacao": forma}, timeout=30)
            r.raise_for_status()
            self._ids[forma] = {p["codigoprocedimento"]: p["procedimento_id"] for p in r.json()}
        try:
            return self._ids[forma][codigo]
        except KeyError:
            raise RuntimeError(f"procedimento {codigo} não existe mais na forma {forma} no PPI") from None

    def consultar(self, subgrupo, forma="", procedimento=""):
        """Consulta com até 3 tentativas: o site do PPI às vezes não responde."""
        for tentativa in range(1, 4):
            try:
                return self._consultar(subgrupo, forma, procedimento)
            except requests.RequestException:
                if tentativa == 3:
                    raise RuntimeError("o site do PPI não respondeu (3 tentativas) — tente de novo mais tarde") from None
                time.sleep(5 * tentativa)

    def _consultar(self, subgrupo, forma="", procedimento=""):
        payload = {
            "_token": self.g.obter_csrf_token(self.s), "_method": "POST",
            "competencia": self.competencia,
            "foco-pesquisa": "2",            # Atendimento
            "discriminar-por": "1",          # Gestão/Município Origem
            "tipo-origem": "", "origem": "", "origem-id": "",
            "atendimento": self.atendimento, "atendimento-id": str(self.atendimento_id),
            "grupo": "02", "subgrupo": subgrupo, "forma-organizacao": forma,
            "procedimento": self.id_procedimento(forma, procedimento) if procedimento else "",
            "complexidade": "", "sistema-informacao": "", "origem-recurso": "", "gestao-recurso": "",
        }
        r = self.s.post(self.base + "consulta-ppi", data=payload, timeout=90)
        r.raise_for_status()
        r.encoding = "utf-8"
        time.sleep(0.6)  # gentileza com o servidor
        return parse_por_origem(r.text), total_oficial(r.text)


# --- planilha -------------------------------------------------------------------

def _limpar(ws, de_linha=1):
    for row in ws.iter_rows(min_row=de_linha, max_row=max(ws.max_row, de_linha)):
        for c in row:
            c.value = None


def _aba_lista(ws, linhas, titulo_qtd="Quantidades Anuais"):
    """Abas por procedimento (e SUS AC): município, qtd, valor, SADT + TOTAL RECURSO."""
    _limpar(ws, 2)
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "Município de Origem", titulo_qtd, "Valores Anuais", "Valor SADT"
    for i, (m, q, v, s) in enumerate(linhas, start=2):
        ws.cell(i, 1, m); ws.cell(i, 2, q); ws.cell(i, 3, v); ws.cell(i, 4, s)
    n = len(linhas) + 1
    t = n + 1
    ws.cell(t, 1, "TOTAL RECURSO"); ws.cell(t, 2, f"=SUM(B2:B{n})"); ws.cell(t, 3, f"=SUM(C2:C{n})"); ws.cell(t, 4, 0)


def _copiar_estilo(ws, origem, destino):
    for col in range(1, ws.max_column + 1):
        a, b = ws.cell(origem, col), ws.cell(destino, col)
        if a.has_style:
            b._style = a._style


# --- grifos e observações do mês anterior -----------------------------------------

def _tem_cor(c):
    try:
        return (c.fill is not None and c.fill.fill_type == "solid"
                and isinstance(c.fill.fgColor.rgb, str) and c.fill.fgColor.rgb not in ("00000000", "FFFFFFFF"))
    except (AttributeError, TypeError):
        return False


def _nome_municipio(txt):
    """'310080 - Aguanil' -> 'aguanil'; 'Aguanil SUS AP' fica como veio."""
    txt = str(txt or "").strip()
    return (txt.split(" - ", 1)[1] if " - " in txt else txt).strip().lower()


def _ler_anotacoes(wb):
    """{aba: {municipio: {"cores": {col: estilo}, "extras": {col: (valor, estilo)}}}} do modelo.

    SUS AP / SUS COLPO têm o nome em fórmula: usamos o município da mesma linha da Geral0203.
    """
    geral = wb["Geral0203"]
    out = {}
    for aba, ncols in (("SUS AP", 4), ("SUS COLPO", 4), ("SUS AC", 4)):
        ws = wb[aba]
        anot = {}
        for r in range(2, ws.max_row + 1):
            chave = _nome_municipio(geral.cell(r, 1).value if aba != "SUS AC" else ws.cell(r, 1).value)
            if not chave or chave.startswith("total"):
                continue
            cores = {c: ws.cell(r, c)._style for c in range(1, ncols + 1) if _tem_cor(ws.cell(r, c))}
            extras = {}
            for c in range(ncols + 1, ws.max_column + 1):
                cel = ws.cell(r, c)
                if cel.value is not None and not (isinstance(cel.value, str) and cel.value.startswith("=")):
                    extras[c] = (cel.value, cel._style)
                elif _tem_cor(cel):
                    extras[c] = (None, cel._style)
            if cores or extras:
                anot[chave] = {"cores": cores, "extras": extras}
        out[aba] = anot
    return out


def _limpar_anotacoes(wb):
    """Tira grifos e observações das posições antigas (as linhas mudam de mês para mês)."""
    for aba in ("SUS AP", "SUS COLPO", "SUS AC"):
        ws = wb[aba]
        vazio = ws.cell(1, ws.max_column + 5)._style  # célula nunca usada = estilo padrão
        limpo = {c: (ws.cell(2, c)._style if not _tem_cor(ws.cell(2, c)) else vazio) for c in range(1, 5)}
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cel = ws.cell(r, c)
                if c > 4:
                    cel.value = None
                    cel._style = vazio
                elif _tem_cor(cel):
                    cel._style = limpo[c]


def _aplicar_anotacoes(wb, anotacoes, municipios_por_aba):
    """Reaplica grifos/observações na linha do mesmo município. Devolve os que sumiram."""
    perdidos = []
    for aba, anot in anotacoes.items():
        ws = wb[aba]
        linhas = {_nome_municipio(m): r for r, m in municipios_por_aba[aba]}
        for chave, a in anot.items():
            r = linhas.get(chave)
            if r is None:
                perdidos.append(f"{aba}: {chave}")
                continue
            for c, estilo in a["cores"].items():
                ws.cell(r, c)._style = estilo
            for c, (valor, estilo) in a["extras"].items():
                ws.cell(r, c).value = valor
                ws.cell(r, c)._style = estilo
    return perdidos


def montar_planilha(modelo, destino, geral, procs, sus_ac):
    """geral: linhas do subgrupo 0203; procs: {codigo: linhas}; sus_ac: linhas do 0202."""
    wb = openpyxl.load_workbook(modelo)
    anotacoes = _ler_anotacoes(wb)
    _limpar_anotacoes(wb)
    for _, cod in PROCEDIMENTOS:
        _aba_lista(wb[cod], procs[cod])

    # SUS AC: mantém a coluna D (Valor SHIFT = valor anual / 12)
    ws = wb["SUS AC"]
    _limpar(ws, 2)
    for i, (m, q, v, _s) in enumerate(sus_ac, start=2):
        ws.cell(i, 1, m); ws.cell(i, 2, q); ws.cell(i, 3, v); ws.cell(i, 4, f"=(C{i}/12)")
    n = len(sus_ac) + 1
    ws.cell(n + 1, 1, "TOTAL RECURSO"); ws.cell(n + 1, 2, f"=SUM(B2:B{n})"); ws.cell(n + 1, 3, f"=SUM(C2:C{n})"); ws.cell(n + 1, 4, 0)

    # Geral0203: total do subgrupo por município + XLOOKUP em cada aba de procedimento
    ws = wb["Geral0203"]
    _limpar(ws, 2)
    ws["A1"], ws["B1"] = "Município de Origem", "Valores Anuais"
    for j, (_, cod) in enumerate(PROCEDIMENTOS, start=3):
        ws.cell(1, j, int(cod))
    letras = "CDEFGHI"
    for i, (m, _q, v, _s) in enumerate(geral, start=2):
        _copiar_estilo(ws, 2, i)
        ws.cell(i, 1, m); ws.cell(i, 2, v)
        for j, (_, cod) in enumerate(PROCEDIMENTOS):
            ws[f"{letras[j]}{i}"] = f"=_xlfn.XLOOKUP(A{i},'{cod}'!A:A,'{cod}'!C:C,0,0)"
    n = len(geral) + 1
    ws.cell(n + 1, 1, "Total Recurso"); ws.cell(n + 1, 2, f"=SUM(B2:B{n})")
    _copiar_estilo(ws, 2, n + 1)
    for letra in letras:
        ws[f"{letra}{n + 1}"] = f"=SUM({letra}2:{letra}{n})"
    ws.cell(n + 3, 1, FONTE)

    # SUS AP / SUS COLPO: uma linha por município da Geral0203
    for nome, sufixo in (("SUS AP", "SUS AP"), ("SUS COLPO", "SUS COLPO")):
        ws = wb[nome]
        _limpar(ws, 2)
        shift = "ShiftAP" if nome == "SUS AP" else "ShiftColpo"
        for i in range(2, n + 1):
            _copiar_estilo(ws, 2, i)
            ws[f"A{i}"] = f'=RIGHT(Geral0203!A{i},LEN(Geral0203!A{i})-SEARCH(" - ",Geral0203!A{i})-2)&" {sufixo}"'
            if nome == "SUS AP":
                ws[f"B{i}"] = f"=(Geral0203!B{i}-SUM(Geral0203!C{i}:I{i}))/12"
                ws[f"D{i}"] = f"=C{i}-B{i}"
            else:
                ws[f"B{i}"] = f"=SUM(Geral0203!C{i}:I{i})/12"
                ws[f"D{i}"] = f"=B{i}-C{i}"
            ws[f"C{i}"] = f"=_xlfn.XLOOKUP(A{i},{shift}!C:C,{shift}!B:B,0,0)"
        if ws.auto_filter.ref:
            ws.auto_filter.ref = f"A1:D{n}"

    municipios = {
        "SUS AP": [(i, m) for i, (m, *_r) in enumerate(geral, start=2)],
        "SUS COLPO": [(i, m) for i, (m, *_r) in enumerate(geral, start=2)],
        "SUS AC": [(i, m) for i, (m, *_r) in enumerate(sus_ac, start=2)],
    }
    montar_planilha.perdidos = _aplicar_anotacoes(wb, anotacoes, municipios)
    wb.calculation.fullCalcOnLoad = True  # Excel recalcula as fórmulas ao abrir
    # Abre sempre na primeira aba, no topo (o modelo guarda a última posição usada)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
        for sv in ws.views.sheetView:
            sv.topLeftCell = "A1"
            if sv.pane is not None and sv.pane.state != "frozen":
                sv.pane = None  # painel dividido rolado: volta à vista simples
            if sv.pane is None:
                sv.selection = [Selection(activeCell="A1", sqref="A1")]
    wb.active = 0
    for v in wb.views:
        v.firstSheet = 0
        v.activeTab = 0
    wb.worksheets[0].sheet_view.tabSelected = True
    destino.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(destino)
        # grava o resultado das fórmulas: o LibreOffice não recalcula ao abrir
        xlsx_calc.gravar_valores(destino)
    except PermissionError:
        raise RuntimeError(f"feche a planilha {destino.name} no Excel e gere de novo") from None


# --- tela / trabalho ------------------------------------------------------------

def situacao():
    PASTA.mkdir(parents=True, exist_ok=True)
    gerados = sorted(SAIDA.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True) if SAIDA.is_dir() else []
    return {
        "tem_modelo": MODELO.is_file(),
        "modelo_em": MODELO.stat().st_mtime if MODELO.is_file() else None,
        "gerados": [{"nome": p.name, "em": p.stat().st_mtime} for p in gerados[:12]],
    }


def salvar_modelo(arquivo_storage):
    PASTA.mkdir(parents=True, exist_ok=True)
    tmp = PASTA / "modelo_novo.xlsx"
    arquivo_storage.save(tmp)
    try:
        # .xlsx corrompido ou sem as abas: o modelo atual fica como está
        wb = openpyxl.load_workbook(tmp, read_only=True)
        try:
            faltam = [n for n in ["Geral0203", "SUS AP", "SUS COLPO", "SUS AC", "ShiftAP", "ShiftColpo"]
                      + [c for _, c in PROCEDIMENTOS] if n not in wb.sheetnames]
        finally:
            wb.close()
        if faltam:
            raise ValueError("a planilha não tem as abas: " + ", ".join(faltam))
        shutil.move(tmp, MODELO)
    finally:
        tmp.unlink(missing_ok=True)


def nome_arquivo(competencia):
    return f"TETO PPI - {MESES[int(competencia[4:]) - 1]} {competencia[:4]}.xlsx"


def executar(job):
    competencia = job.params["competencia"]
    if not MODELO.is_file():
        raise RuntimeError("envie a planilha-modelo (a do mês anterior) antes de gerar")
    etapas = ["Subgrupo 0203", "Subgrupo 0202"] + [c for _, c in PROCEDIMENTOS] + ["Planilha"]
    job.definir_itens(etapas)
    job.progresso(0, len(etapas))
    job.escrever(f"Competência {ppi_engine.formatar_competencia(competencia)} · Atendimento {ATENDIMENTO}")
    c = Consulta(competencia)
    feito = 0

    def etapa(nome, fn):
        nonlocal feito
        if job.cancelado:
            raise RuntimeError("cancelado")
        job.marcar(nome, "rodando")
        t0 = time.monotonic()
        try:
            res = fn()
        except Exception:
            job.marcar(nome, "erro", round(time.monotonic() - t0, 1))
            raise
        job.marcar(nome, "ok", round(time.monotonic() - t0, 1))
        feito += 1
        job.progresso(feito, len(etapas))
        return res

    conferencia = []

    def consultar_e_conferir(nome, *args):
        linhas, total = c.consultar(*args)
        try:
            qtd, val = conferir(nome, linhas, total)
        except RuntimeError as e:
            job.escrever("✘ " + str(e))
            raise
        conferencia.append({"etapa": nome, "municipios": len(linhas), "qtd": qtd, "valor": round(val, 2)})
        job.escrever(f"✔ {nome}: {len(linhas)} municípios · qtd {qtd:,.0f} · R$ {val:,.2f} — confere com o Total Recurso do site"
                     .replace(",", "X").replace(".", ",").replace("X", "."))
        for rot, q, v in total["outros"]:
            job.escrever(f"   (fora da planilha: bloco {rot} — qtd {q:,.0f} · R$ {v:,.2f})"
                         .replace(",", "X").replace(".", ",").replace("X", "."))
        return linhas

    geral = etapa("Subgrupo 0203", lambda: consultar_e_conferir("Subgrupo 0203", "0203"))
    if not geral:
        raise RuntimeError("o PPI não devolveu nenhuma linha para o 0203 — confira a competência")
    sus_ac = etapa("Subgrupo 0202", lambda: consultar_e_conferir("Subgrupo 0202", "0202"))
    procs = {}
    for forma, cod in PROCEDIMENTOS:
        procs[cod] = etapa(cod, lambda f=forma, k=cod: consultar_e_conferir(k, "0203", f, k))

    destino = SAIDA / nome_arquivo(competencia)
    etapa("Planilha", lambda: montar_planilha(MODELO, destino, geral, procs, sus_ac))
    job.escrever(f"Planilha salva: {destino.name}")
    for p in getattr(montar_planilha, "perdidos", []):
        job.escrever(f"⚠ observação do mês anterior não aplicada (município saiu do PPI): {p}")
    return {"arquivo": destino.name, "conferencia": conferencia, "gerado_em": datetime.now().strftime("%d/%m/%Y %H:%M")}

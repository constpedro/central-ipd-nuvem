"""
Os 853 municípios de MG com nome oficial e código IBGE (planilha de regionalização
da SES-MG), para achar a cidade certa mesmo quando o nome vem escrito diferente:
sem acento, com parêntese ("Paraíso (São Sebastião do Paraíso)"), com hífen ou
apóstrofo trocado, ou com um erro pequeno de digitação.

Guardado em data/municipios_mg.json (baixado uma vez da SES-MG).
"""
import difflib
import json
import re
import unicodedata
from io import BytesIO

import requests

import config

ARQ = config.DATA_DIR / "municipios_mg.json"
URL_COORD = "https://raw.githubusercontent.com/kelvins/municipios-brasileiros/main/csv/municipios.csv"
URL_SES = "https://www.saude.mg.gov.br/wp-content/uploads/2026/02/Planilha-de-Regionalizacao_Versao-2026.xlsx"
_cache = None


def chave(nome):
    """'Pingo-d'Água' e 'pingo dagua' viram a mesma chave."""
    s = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]", "", s)


def baixar():
    """Lê a planilha da SES-MG e grava data/municipios_mg.json."""
    from openpyxl import load_workbook
    r = requests.get(URL_SES, timeout=60, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    ws = load_workbook(BytesIO(r.content), read_only=True).worksheets[0]
    lista = []
    for linha in ws.iter_rows(values_only=True):
        if linha and isinstance(linha[1], (int, float)) and str(int(linha[1])).startswith("31"):
            lista.append({"codigo": str(int(linha[1])), "nome": str(linha[2]).strip(),
                          "populacao": linha[3], "regional": linha[4], "microrregiao": linha[6],
                          "macrorregiao": linha[8]})
    if len(lista) < 800:
        raise RuntimeError(f"a planilha da SES-MG veio com só {len(lista)} municípios")
    _com_coordenadas(lista)
    ARQ.write_text(json.dumps(lista, ensure_ascii=False, indent=1), encoding="utf-8")
    return lista


def _com_coordenadas(lista):
    """Latitude, longitude e DDD (base aberta kelvins/municipios-brasileiros, código IBGE de 7 dígitos)."""
    import csv
    r = requests.get(URL_COORD, timeout=60)
    r.raise_for_status()
    coord = {l["codigo_ibge"][:6]: l for l in csv.DictReader(r.text.splitlines())}
    for m in lista:
        c = coord.get(m["codigo"])
        if c:
            m.update(lat=float(c["latitude"]), lon=float(c["longitude"]), ddd=c["ddd"])


def todos():
    global _cache
    if _cache is None:
        try:
            _cache = json.loads(ARQ.read_text(encoding="utf-8"))
            if _cache and "lat" not in _cache[0]:  # arquivo antigo, sem coordenadas
                _com_coordenadas(_cache)
                ARQ.write_text(json.dumps(_cache, ensure_ascii=False, indent=1), encoding="utf-8")
        except (OSError, ValueError):
            try:
                _cache = baixar()
            except Exception:  # noqa: BLE001 — sem a planilha, segue sem a conferência
                _cache = []
    return _cache


def resolver(nome):
    """{"codigo", "nome", ...} do município de MG, ou None se não achar com segurança."""
    lista = todos()
    if not lista:
        return None
    por_chave = {chave(m["nome"]): m for m in lista}
    tentativas = [nome]
    dentro = re.search(r"\(([^)]+)\)", nome)
    if dentro:  # "Paraíso (São Sebastião do Paraíso)": o de dentro e o de fora
        tentativas += [dentro.group(1), re.sub(r"\([^)]*\)", "", nome)]
    tentativas = [re.sub(r"\s*[-/]\s*MG\s*$", "", t, flags=re.I).strip() for t in tentativas]
    for t in tentativas:
        if chave(t) in por_chave:
            return por_chave[chave(t)]
    for t in tentativas:  # erro pequeno de digitação ("Sao Tome das Letra")
        perto = difflib.get_close_matches(chave(t), por_chave, n=1, cutoff=0.88)
        if perto:
            return por_chave[perto[0]]
    return None


# Nomes com que o site do PPI cadastrou a cidade, quando diferem do oficial da SES-MG.
NOMES_NO_PPI = {"316520": ["São Thomé das Letras"]}


def nomes_para_busca(codigo, nome):
    """O nome oficial e as grafias alternativas conhecidas do PPI."""
    return [nome] + [n for n in NOMES_NO_PPI.get(codigo, []) if n != nome]


_por_codigo = (None, {})


def por_codigo():
    """{código IBGE: município}; montado uma vez por lista (era refeito a cada distância)."""
    global _por_codigo
    lista = todos()
    if _por_codigo[0] is not lista:
        _por_codigo = (lista, {m["codigo"]: m for m in lista})
    return _por_codigo[1]


def distancia_varginha(codigo):
    """km em linha reta até Varginha, ou None sem coordenadas."""
    from core import contatos
    m = por_codigo().get(codigo) or {}
    return contatos.distancia_varginha(m.get("lat"), m.get("lon"))

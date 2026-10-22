# -*- coding: utf-8 -*-
"""
Tabela de preços dos exames e conferência contra arquivos do Shift.

As tabelas oficiais são três PDFs (Faturado, À vista, Particular), 44 páginas
cada, uma linha por exame no formato

    MNEMÔNICO   DESCRIÇÃO DO EXAME   R$ 123,45

Lemos os três com o pypdf uma vez, juntamos por mnemônico numa tabela única
(mnemônico → descrição + os três preços) e guardamos em data/tabelas.json.
O catálogo do Shift (CSV) complementa com a área laboratorial de cada exame.

A conferência recebe um arquivo qualquer (.csv/.xlsx) que tenha uma coluna de
mnemônico e uma de valor — um export do Shift, uma tabela de consórcio — e
aponta o que está diferente da tabela oficial na modalidade escolhida.
"""
import csv
import io
import json
import re
import unicodedata
from datetime import datetime

import config

RE_LINHA = re.compile(r"^(\S+)\s+(.*?)\s*R\$\s*([\d.]+,\d{2})\s*$")
RE_SO_PRECO = re.compile(r"^R\$\s*([\d.]+,\d{2})\s*$")
# Só pontos, em grupos de 3 depois de um 1º grupo sem zero à esquerda: milhar.
# "12.5", "12.50" e "0.125" continuam decimais.
RE_MILHAR_PONTO = re.compile(r"[+-]?[1-9]\d{0,2}(\.\d{3})+")

_cache = None


def _num(txt):
    return float(txt.replace(".", "").replace(",", "."))


def normalizar(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s).strip().upper()


def nome_arquivo_seguro(nome):
    """Limpa um componente de nome sem destruir letras acentuadas."""
    nome = unicodedata.normalize("NFC", str(nome or ""))
    nome = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", nome)
    nome = re.sub(r"\s+", " ", nome).strip(" .")
    return nome or "sem-nome"


def ler_pdf(caminho):
    """{mnemônico: (descrição, valor)} de um PDF de tabela."""
    from pypdf import PdfReader
    itens = {}
    pendente = ""  # descrição que quebrou de linha (a próxima linha traz só o preço)
    for pagina in PdfReader(str(caminho)).pages:
        for linha in (pagina.extract_text() or "").splitlines():
            linha = linha.strip()
            if not linha or linha.lower().startswith(("mnemônico", "mnemonico")) or ("2026" in linha[:40] and "R$" not in linha):
                continue
            m = RE_LINHA.match(linha)
            if m:
                mnem, desc, valor = m.group(1), m.group(2), _num(m.group(3))
                if pendente:
                    desc = (pendente + " " + linha.rsplit("R$", 1)[0]).strip()
                    mnem = pendente.split()[0]
                    desc = desc[len(mnem):].strip()
                    pendente = ""
                itens[mnem] = (desc.lstrip("_").strip(), valor)
                continue
            m2 = RE_SO_PRECO.match(linha)
            if m2 and pendente:
                partes = pendente.split(None, 1)
                itens[partes[0]] = ((partes[1] if len(partes) > 1 else "").lstrip("_").strip(), _num(m2.group(1)))
                pendente = ""
                continue
            # linha sem preço: guarda para juntar com a próxima
            pendente = (pendente + " " + linha).strip() if pendente else linha
    return itens


def ler_csv_shift(caminho):
    """{mnemônico: (descrição, valor, grupo)} de uma exportação de procedimentos do
    Shift (preço em "Quantidade") ou das tabelas enxutas de out/2026 (Mnemônico;
    Descrição; Valor; Grupo)."""
    saida = {}
    with open(caminho, encoding="latin-1", newline="") as f:
        for row in csv.DictReader(f, delimiter=";"):
            mnem = (row.get("Mnemônico") or "").strip()
            valor = para_numero(row.get("Valor") if row.get("Valor") is not None else row.get("Quantidade"))
            valor = None if valor is None else round(valor, 2)
            grupo = (row.get("Grupo") or "").strip()
            if mnem and valor is not None:
                saida[mnem] = ((row.get("Descrição") or "").strip().lstrip("_").strip(), valor,
                               "" if grupo == "0" else grupo)
    return saida


def ler_tabela(caminho):
    return ler_csv_shift(caminho) if caminho.suffix.lower() == ".csv" else ler_pdf(caminho)


def ler_catalogo_shift():
    """{mnemônico: {descricao, area}} do CSV de procedimentos do Shift."""
    p = config.SHIFT_PROCEDIMENTOS_CSV
    if not p.is_file():
        return {}
    saida = {}
    with open(p, encoding="latin-1", newline="") as f:
        for row in csv.DictReader(f, delimiter=";"):
            mnem = (row.get("Mnemônico") or "").strip()
            if mnem:
                saida[mnem] = {
                    "descricao": (row.get("Descrição") or "").strip(),
                    "area": (row.get("Área Laboratorial Padrão") or "").strip(),
                }
    return saida


def _texto(v):
    return "" if v is None else str(v).strip()


def _codigo(v):
    """Conserva códigos numéricos de planilhas sem o sufixo .0."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return _texto(v)


def _coluna(headers, *termos):
    for i, header in enumerate(headers):
        n = normalizar(header)
        if any(termo in n for termo in termos):
            return i
    return None


def _ler_extra_xlsx(caminho, chave):
    """Lê uma tabela adicional em XLSX para o formato comum da consulta."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    linhas = [list(row) for row in ws.iter_rows(values_only=True)]
    cabecalho = None
    inicio = 0
    for i, row in enumerate(linhas):
        norm = [normalizar(v) for v in row]
        tem_mnemonico = any("MNEMONICO" in v for v in norm)
        tem_cislav = any(v.startswith("ITEM") for v in norm) and any(v.startswith("QUANT") for v in norm)
        if tem_mnemonico or tem_cislav:
            if tem_cislav and i + 1 < len(linhas):
                baixo = linhas[i + 1]
                n = max(len(row), len(baixo))
                cabecalho = [" / ".join(_texto(x) for x in (row[j] if j < len(row) else None,
                                                              baixo[j] if j < len(baixo) else None) if _texto(x))
                             for j in range(n)]
                inicio = i + 2
            else:
                cabecalho = [_texto(x) for x in row]
                inicio = i + 1
            break
    if cabecalho is None:
        raise ValueError("cabeçalho não reconhecido")

    if chave == "cislav":
        idx_mn = None
        idx_desc = _coluna(cabecalho, "PROCEDIMENTO")
    else:
        idx_mn = _coluna(cabecalho, "MNEMONICO")
        idx_desc = _coluna(cabecalho, "NOME", "DESCRICAO", "EXAME", "PROCEDIMENTO")
    idx_code = _coluna(cabecalho, "CODIGO SUS", "COD SUS")
    idx_valor = _coluna(cabecalho, "TABELA IPD", "VALOR INICIAL", "VALOR MINIMO", "VALOR", "PRECO")
    idx_minimo = _coluna(cabecalho, "VALOR MINIMO")
    idx_qtd = _coluna(cabecalho, "QUANT")
    idx_item = _coluna(cabecalho, "ITEM")
    if idx_desc is None or idx_valor is None:
        raise ValueError("colunas de descrição e valor não reconhecidas")

    itens = []
    for row in linhas[inicio:]:
        def valor_col(idx):
            return row[idx] if idx is not None and idx < len(row) else None

        mn = _texto(valor_col(idx_mn))
        if normalizar(mn) in {"NAO", "N/A", "NAO SE APLICA", "JUNTO COM CULTURA"}:
            mn = ""
        descricao = _texto(valor_col(idx_desc))
        codigo_sus = _codigo(valor_col(idx_code))
        identificador = mn or codigo_sus or _codigo(valor_col(idx_item))
        valor = para_numero(valor_col(idx_valor))
        if not identificador or not descricao or valor is None:
            continue
        itens.append({
            "identificador": identificador,
            "mnemonico": mn,
            "descricao": descricao,
            "codigo_sus": codigo_sus,
            "codigo_tuss": "",
            "quantidade": para_numero(valor_col(idx_qtd)),
            "valor": valor,
            "valor_minimo": para_numero(valor_col(idx_minimo)),
        })
    return itens


def _ler_extra_docx(caminho):
    """Lê a referência SUS/TUSS, que não possui coluna de preço."""
    from docx import Document

    doc = Document(caminho)
    itens = []
    for tabela in doc.tables:
        if not tabela.rows:
            continue
        primeira = [_texto(c.text) for c in tabela.rows[0].cells]
        norm = [normalizar(v) for v in primeira]
        tem_cabecalho = any("MNEMONICO" in v for v in norm)
        if tem_cabecalho:
            headers = norm
            inicio = 1
            idx_mn = _coluna(headers, "MNEMONICO")
            idx_desc = _coluna(headers, "EXAME", "DESCRICAO", "NOME")
            idx_sus = _coluna(headers, "CODIGO SUS", "COD SUS")
            idx_tuss = _coluna(headers, "CODIGO TUSS", "COD TUSS")
        else:
            # A última tabela é uma continuação sem cabeçalho.
            inicio = 0
            idx_mn, idx_desc, idx_sus, idx_tuss = 0, 1, 2, None
        if idx_mn is None or idx_desc is None:
            continue
        for row in tabela.rows[inicio:]:
            valores = [_texto(c.text) for c in row.cells]
            mn = valores[idx_mn] if idx_mn < len(valores) else ""
            descricao = valores[idx_desc] if idx_desc < len(valores) else ""
            codigo_sus = valores[idx_sus] if idx_sus is not None and idx_sus < len(valores) else ""
            codigo_tuss = valores[idx_tuss] if idx_tuss is not None and idx_tuss < len(valores) else ""
            if not mn or not descricao:
                continue
            itens.append({
                "identificador": mn,
                "mnemonico": mn,
                "descricao": descricao,
                "codigo_sus": codigo_sus,
                "codigo_tuss": codigo_tuss,
                "quantidade": None,
                "valor": None,
            })
    return itens


def reimportar():
    """Relê os PDFs e o CSV e regrava data/tabelas.json."""
    global _cache
    por_modalidade = {}
    fontes = {}
    for mod, caminho in config.TABELAS_PRINCIPAIS.items():
        if caminho.is_file():
            por_modalidade[mod] = ler_tabela(caminho)
            fontes[mod] = {"arquivo": caminho.name, "itens": len(por_modalidade[mod])}
        else:
            por_modalidade[mod] = {}
            fontes[mod] = {"arquivo": caminho.name, "itens": 0, "erro": "arquivo não encontrado"}
    extras = {}
    fontes_adicionais = {}
    for chave, fonte in config.TABELAS_ADICIONAIS.items():
        caminho = fonte["arquivo"]
        info = {"arquivo": caminho.name, "itens": 0}
        if not caminho.is_file():
            info["erro"] = "arquivo não encontrado"
            extras[chave] = []
        else:
            try:
                if fonte["tipo"] == "xlsx":
                    extra = _ler_extra_xlsx(caminho, chave)
                elif fonte["tipo"] == "docx":
                    extra = _ler_extra_docx(caminho)
                else:
                    raise ValueError("tipo de arquivo não suportado")
                extras[chave] = extra
                info["itens"] = len(extra)
            except Exception as exc:  # noqa: BLE001
                extras[chave] = []
                info["erro"] = str(exc)
        fontes_adicionais[chave] = info

    shift = ler_catalogo_shift()
    todos = set()
    for d in por_modalidade.values():
        todos |= set(d)
    itens = []
    for mnem in sorted(todos):
        desc = ""
        for mod in por_modalidade:
            if mnem in por_modalidade[mod]:
                desc = por_modalidade[mod][mnem][0]
                break
        grupo = next((d[mnem][2] for d in por_modalidade.values() if len(d.get(mnem, ())) > 2 and d[mnem][2]), "")
        item = {"mnemonico": mnem, "descricao": desc, "area": shift.get(mnem, {}).get("area", ""), "grupo": grupo}
        for mod, tabela in por_modalidade.items():
            item[mod] = tabela.get(mnem, (None, None))[1]
        itens.append(item)
    _cache = {
        "importado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "fontes": fontes,
        "fontes_adicionais": fontes_adicionais,
        "extras": extras,
        "n_shift": len(shift),
        "itens": itens,
    }
    config.TABELAS_CACHE.write_text(json.dumps(_cache, ensure_ascii=False), encoding="utf-8")
    return _cache


def carregar():
    """Lê data/tabelas.json (em memória depois da 1ª vez). Nunca reimporta aqui:
    a reimportação leva ~1 min e roda na fila de jobs (app._importar_tabelas_job)."""
    global _cache
    if _cache is None:
        if not config.TABELAS_CACHE.is_file():
            raise FileNotFoundError("tabela ainda não importada — use 'Reimportar tabelas'")
        _cache = json.loads(config.TABELAS_CACHE.read_text(encoding="utf-8"))
    return _cache


def cache_desatualizado():
    """True quando o JSON foi gerado com outro conjunto de TABELAS_ADICIONAIS
    (alguém mexeu no config.py) e vale reimportar."""
    try:
        c = carregar()
        principais = {m: c.get("fontes", {}).get(m, {}).get("arquivo") for m in config.TABELAS_PRINCIPAIS}
        return (set(c.get("extras", {})) != set(config.TABELAS_ADICIONAIS)
                or principais != {m: p.name for m, p in config.TABELAS_PRINCIPAIS.items()})
    except (FileNotFoundError, json.JSONDecodeError):
        return False


_RE_NAO_ALFANUM = re.compile(r"[^A-Z0-9]")
_indice = {"cache": None}  # índice de busca do cache atual (refeito quando ele muda)


def _indices():
    """Textos já normalizados para a busca, calculados uma vez por importação.

    Antes cada tecla digitada na busca normalizava de novo os ~5 mil textos das
    tabelas (~150 ms por busca). O índice acompanha o objeto do cache: quando
    reimportar() troca o cache, ele é refeito sozinho.
    """
    cache = carregar()
    if _indice["cache"] is cache:
        return _indice
    principais = [(it, it["mnemonico"].upper(), normalizar(it["descricao"])) for it in cache["itens"]]
    por_mnemonico = {}
    for it, mn, _ in principais:
        por_mnemonico.setdefault(mn, it)

    # Preços das tabelas adicionais para casar com a tabela principal. A CISSM
    # traz mnemônico; a CISLAV identifica o exame pelo código SUS e pela
    # descrição. Por isso dois caminhos: {chave: (por_mnemonico, por_descricao)}.
    adicionais = {}
    extras = []  # (ordem da fonte, nome da fonte, item, campos normalizados, compactos, descrição)
    for ordem, (chave, fonte) in enumerate(config.TABELAS_ADICIONAIS.items()):
        por_mn, por_desc = {}, {}
        for item in cache.get("extras", {}).get(chave, []):
            descricao = normalizar(item.get("descricao") or "")
            campos = [normalizar(item.get(c) or "") for c in ("identificador", "mnemonico", "codigo_sus", "codigo_tuss")]
            extras.append((ordem, (chave, fonte["nome"]), item, campos,
                           [_RE_NAO_ALFANUM.sub("", v) for v in campos], descricao))
            valor = item.get("valor")
            if valor is None:
                continue
            for campo in ("mnemonico", "identificador"):
                mn = (item.get(campo) or "").strip().upper()
                if mn:
                    por_mn.setdefault(mn, valor)
            if descricao:
                por_desc.setdefault(descricao, valor)
        adicionais[chave] = (por_mn, por_desc)

    _indice.update(cache=cache, principais=principais, por_mnemonico=por_mnemonico,
                   adicionais=adicionais, extras=extras)
    return _indice


def _com_adicionais(itens):
    """Devolve cópias dos itens com o valor de cada tabela adicional junto."""
    indice = _indices()["adicionais"]
    saida = []
    for it in itens:
        copia = dict(it)
        mn = (copia.get("mnemonico") or "").strip().upper()
        descricao = normalizar(copia.get("descricao") or "")
        for chave, (por_mnemonico, por_descricao) in indice.items():
            copia[chave] = por_mnemonico.get(mn)
            if copia[chave] is None and descricao:
                copia[chave] = por_descricao.get(descricao)
        saida.append(copia)
    return saida


def buscar(q, limite=60):
    """Busca por mnemônico (prefixo tem prioridade) ou por palavras da descrição."""
    q = normalizar(q)
    if not q:
        return []
    palavras = q.split()
    exatos, prefixo, resto = [], [], []
    for it, mn, desc in _indices()["principais"]:
        if mn == q:
            exatos.append(it)
        elif mn.startswith(q):
            prefixo.append(it)
        elif all(p in desc or p in mn for p in palavras):
            resto.append(it)
    return _com_adicionais((exatos + prefixo + resto)[:limite])


def buscar_adicionais(q, limite=80):
    """Busca nas fontes adicionais por mnemônico, código ou descrição."""
    q = normalizar(q)
    if not q:
        return []
    compacto_q = _RE_NAO_ALFANUM.sub("", q)
    palavras = q.split()
    achados = []
    for ordem, nome_fonte, item, campos_norm, campos_compactos, descricao in _indices()["extras"]:
        if compacto_q and compacto_q in campos_compactos:
            score = 0
        elif any(v.startswith(q) for v in campos_norm if v):
            score = 1
        elif any(compacto_q and compacto_q in v for v in campos_compactos):
            score = 1
        elif all(p in descricao or any(p in v for v in campos_norm) for p in palavras):
            score = 2
        else:
            continue
        achados.append((score, ordem, descricao, item, nome_fonte))
    achados.sort(key=lambda x: x[:3])
    return [dict(item, chave=fonte[0], fonte=fonte[1]) for _, _, _, item, fonte in achados[:limite]]


def obter(mnemonico):
    return _indices()["por_mnemonico"].get((mnemonico or "").strip().upper())


# --- conferência ----------------------------------------------------------------

def ler_arquivo_tabular(nome, conteudo):
    """Devolve (cabecalho, linhas) de um .csv (; ou ,) ou .xlsx enviado."""
    nome = nome.lower()
    if nome.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        linhas = [["" if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        texto = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                texto = conteudo.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        primeira = texto.splitlines()[0] if texto else ""
        if primeira.lower().startswith("sep="):
            partes = texto.split("\n", 1)
            texto = partes[1] if len(partes) > 1 else ""
            primeira = texto.splitlines()[0] if texto else ""
        delim = ";" if primeira.count(";") >= primeira.count(",") else ","
        linhas = list(csv.reader(io.StringIO(texto), delimiter=delim))
    # cabeçalho = primeira linha com 2+ células preenchidas
    for i, l in enumerate(linhas):
        if sum(1 for c in l if str(c).strip()) >= 2:
            return [str(c).strip() for c in l], [r for r in linhas[i + 1:] if any(str(c).strip() for c in r)]
    return [], []


def detectar_colunas(cabecalho):
    """Índices sugeridos para mnemônico e valor a partir dos nomes de coluna."""
    def acha(*chaves):
        for i, c in enumerate(cabecalho):
            n = normalizar(c)
            if any(k in n for k in chaves):
                return i
        return None
    col_mnem = acha("MNEM")
    col_valor = acha("VALOR", "PRECO", "VLR", "TABELA IPD")
    col_desc = acha("DESCRI", "NOME", "EXAME")
    return col_mnem, col_valor, col_desc


def para_numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("R$", "").replace(" ", "").replace("\xa0", "")
    if not s:
        return None
    if "," in s and "." in s:
        # o separador decimal é o que aparece por último: "1.234,56" e "1,234.56"
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif RE_MILHAR_PONTO.fullmatch(s):
        s = s.replace(".", "")  # "1.234" e "1.500.000": ponto de milhar, como se escreve no Brasil
    try:
        return float(s)
    except ValueError:
        return None


def conferir(cabecalho, linhas, col_mnem, col_valor, modalidade, col_desc=None, tolerancia=0.005):
    """Compara cada linha do arquivo com a tabela oficial na modalidade."""
    tabela = _indices()["por_mnemonico"]
    divergentes, iguais, sem_tabela, sem_valor = [], [], [], []
    for r in linhas:
        if col_mnem >= len(r):
            continue
        mnem = str(r[col_mnem]).strip().upper()
        if not mnem:
            continue
        valor = para_numero(r[col_valor]) if col_valor is not None and col_valor < len(r) else None
        desc = str(r[col_desc]).strip() if col_desc is not None and col_desc < len(r) else ""
        it = tabela.get(mnem)
        if not it:
            sem_tabela.append({"mnemonico": mnem, "descricao": desc, "valor_arquivo": valor})
            continue
        oficial = it.get(modalidade)
        reg = {"mnemonico": mnem, "descricao": it["descricao"] or desc, "valor_arquivo": valor, "valor_tabela": oficial}
        if valor is None:
            sem_valor.append(reg)
        elif oficial is None:
            reg["diferenca"] = None
            divergentes.append(reg)
        elif abs(valor - oficial) > tolerancia:
            reg["diferenca"] = round(valor - oficial, 2)
            reg["percentual"] = round(100 * (valor - oficial) / oficial, 1) if oficial else None
            divergentes.append(reg)
        else:
            iguais.append(reg)
    divergentes.sort(key=lambda x: -abs(x["diferenca"] or 0))
    return {
        "modalidade": modalidade,
        "modalidade_nome": config.MODALIDADES.get(modalidade, modalidade),
        "n_linhas": len(linhas),
        "divergentes": divergentes,
        "iguais": len(iguais),
        "sem_tabela": sem_tabela,
        "sem_valor": len(sem_valor),
    }


def exportar_conferencia_xlsx(resultado, caminho):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Divergências"
    ws.append(["Mnemônico", "Descrição", "Valor no arquivo", f"Valor na tabela ({resultado['modalidade_nome']})", "Diferença", "%"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for d in resultado["divergentes"]:
        ws.append([d["mnemonico"], d["descricao"], d["valor_arquivo"], d["valor_tabela"], d.get("diferenca"), d.get("percentual")])
    ws2 = wb.create_sheet("Fora da tabela")
    ws2.append(["Mnemônico", "Descrição no arquivo", "Valor no arquivo"])
    for d in resultado["sem_tabela"]:
        ws2.append([d["mnemonico"], d["descricao"], d["valor_arquivo"]])
    for w, col in zip((14, 60, 18, 26, 14, 8), "ABCDEF"):
        ws.column_dimensions[col].width = w
    wb.save(caminho)


GRUPOS_PRODUTO = {"analises_clinicas": "ANALISES CLINICAS", "anatomia_patologica": "ANATOMIA PATOLOGICA",
                  "citologia": "CITOLOGIA"}


def tabela_completa(modalidade, patologia=None, produtos=()):
    """Todos os exames de uma tabela para a proposta "Tabela completa":
    [{mnemonico, descricao, codigo_sus, valor}], com a patologia escolhida por
    cima e filtrados pelo grupo dos produtos marcados (sem produto = todos)."""
    cache = carregar()
    grupos = {GRUPOS_PRODUTO[p] for p in produtos if p in GRUPOS_PRODUTO}
    pat = cache.get("extras", {}).get(config.PATOLOGIAS.get(patologia, ""), [])
    pat_mn = {(x.get("mnemonico") or "").upper(): x for x in pat if x.get("mnemonico")}
    saida = []
    if modalidade in config.MODALIDADES:
        for it in cache["itens"]:
            if grupos and it.get("grupo") and it["grupo"] not in grupos:
                continue
            p = pat_mn.pop(it["mnemonico"].upper(), None)
            valor = p["valor"] if p else it.get(modalidade)
            if valor is not None:
                saida.append({"mnemonico": it["mnemonico"], "descricao": it["descricao"], "codigo_sus": "", "valor": valor})
    else:  # tabela de consórcio (CISLAV/CISSM): sem grupo, vai inteira
        # A CISLAV não traz mnemônico: busca o nosso pela descrição (inteira ou
        # sem o sufixo "– EXAMES ... – GENÉRICO" do consórcio).
        por_desc = {}
        for it in cache["itens"]:
            por_desc.setdefault(normalizar(it["descricao"]), it["mnemonico"])
        for x in cache.get("extras", {}).get(modalidade, []):
            mn = x.get("mnemonico") or ""
            if not mn:
                desc = normalizar(x.get("descricao") or "")
                mn = por_desc.get(desc) or por_desc.get(re.split(r"\s+\W?\s*EXAMES?\b", desc)[0].strip(" -,"), "")
            p = pat_mn.pop(mn.upper(), None) if mn else None
            saida.append({"mnemonico": mn, "descricao": x["descricao"],
                          "codigo_sus": x.get("codigo_sus") or "", "valor": p["valor"] if p else x["valor"]})
    # exames que só existem na patologia (PECASIMP...) entram se a anatomia foi marcada
    if not grupos or "ANATOMIA PATOLOGICA" in grupos:
        saida += [{"mnemonico": x["mnemonico"], "descricao": x["descricao"], "codigo_sus": x.get("codigo_sus") or "",
                   "valor": x["valor"]} for x in pat_mn.values()]
    return saida

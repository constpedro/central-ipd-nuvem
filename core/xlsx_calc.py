# -*- coding: utf-8 -*-
"""
Grava o resultado das fórmulas dentro do .xlsx gerado pelo openpyxl.

O openpyxl salva só a fórmula, sem o valor calculado. O Excel recalcula ao
abrir, mas o LibreOffice (padrão: "nunca recalcular" em .xlsx) mostra as
células vazias. Aqui avaliamos as fórmulas da planilha do Teto PPI e
escrevemos o valor em <v> ao lado de cada <f>, mantendo a fórmula.

Cobre só o que essas planilhas usam: números, textos, + - * /, &, referências
(A1, Aba!A1, 'Aba'!A1), SUM(intervalo), RIGHT, LEN, SEARCH e XLOOKUP de
coluna inteira. Fórmula fora disso fica sem valor (como antes).
"""
import re
import shutil
import tempfile
import zipfile
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import column_index_from_string, range_boundaries

_TOKEN = re.compile(r"""
    \s*(?:
      (?P<num>\d+(?:\.\d+)?)
    | (?P<str>"(?:[^"]|"")*")
    | (?P<ref>(?:(?:'[^']+'|[A-Za-z0-9_]+)!)?\$?[A-Z]{1,3}\$?\d*(?::\$?[A-Z]{1,3}\$?\d*)?)(?![A-Za-z0-9_(])
    | (?P<func>[A-Za-z_.]+)\(
    | (?P<op>[-+*/&(),=<>])
    )""", re.X)


class _Erro(Exception):
    pass


class Calculadora:
    def __init__(self, wb):
        self.wb = wb
        self.cache = {}
        self.pilha = set()

    # -- valores ---------------------------------------------------------------
    def valor(self, aba, ref):
        chave = (aba, ref)
        if chave in self.cache:
            return self.cache[chave]
        if chave in self.pilha:
            raise _Erro("referência circular")
        bruto = self.wb[aba][ref].value
        if isinstance(bruto, str) and bruto.startswith("="):
            self.pilha.add(chave)
            try:
                v = _Parser(self, aba, bruto[1:]).expr_total()
            finally:
                self.pilha.discard(chave)
        else:
            v = bruto
        self.cache[chave] = v
        return v

    def coluna(self, aba, letra):
        ws = self.wb[aba]
        c = column_index_from_string(letra)
        return [self.valor(aba, f"{letra}{r}") for r in range(1, ws.max_row + 1)] if ws.max_column >= c else []

    def intervalo(self, aba, txt):
        c1, r1, c2, r2 = range_boundaries(txt.replace("$", ""))
        ws = self.wb[aba]
        r1, r2 = r1 or 1, r2 or ws.max_row
        out = []
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                out.append(self.valor(aba, ws.cell(r, c).coordinate))
        return out


def _num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return v
    try:
        return float(v)
    except ValueError:
        raise _Erro(f"não é número: {v!r}")


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


class _Parser:
    def __init__(self, calc, aba, formula):
        self.calc, self.aba = calc, aba
        self.toks, pos = [], 0
        formula = formula.strip()
        while pos < len(formula):
            m = _TOKEN.match(formula, pos)
            if not m or m.end() == pos:
                raise _Erro(f"não entendi: {formula[pos:pos + 20]}")
            pos = m.end()
            tipo = m.lastgroup
            self.toks.append((tipo, m.group(tipo)))
        self.i = 0

    def _ver(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def _pegar(self, op=None):
        t = self._ver()
        if op is not None and t != ("op", op):
            raise _Erro(f"esperava {op}")
        self.i += 1
        return t

    def expr_total(self):
        v = self.concat()
        if self.i != len(self.toks):
            raise _Erro("sobrou texto na fórmula")
        return v

    def concat(self):
        v = self.soma()
        while self._ver() == ("op", "&"):
            self._pegar()
            v = _txt(v) + _txt(self.soma())
        return v

    def soma(self):
        v = self.termo()
        while self._ver() in (("op", "+"), ("op", "-")):
            op = self._pegar()[1]
            d = self.termo()
            v = _num(v) + _num(d) if op == "+" else _num(v) - _num(d)
        return v

    def termo(self):
        v = self.fator()
        while self._ver() in (("op", "*"), ("op", "/")):
            op = self._pegar()[1]
            d = _num(self.fator())
            if op == "/" and d == 0:
                raise _Erro("divisão por zero")
            v = _num(v) * d if op == "*" else _num(v) / d
        return v

    def fator(self):
        tipo, txt = self._pegar()
        if tipo == "num":
            return float(txt)
        if tipo == "str":
            return txt[1:-1].replace('""', '"')
        if (tipo, txt) == ("op", "-"):
            return -_num(self.fator())
        if (tipo, txt) == ("op", "("):
            v = self.concat()
            self._pegar(")")
            return v
        if tipo == "ref":
            aba, ref = self._separar(txt)
            if ":" in ref:
                raise _Erro("intervalo fora de função")
            return self.calc.valor(aba, ref.replace("$", ""))
        if tipo == "func":
            return self.funcao(txt.upper().replace("_XLFN.", ""))
        raise _Erro(f"token inesperado {txt}")

    def _separar(self, txt):
        if "!" in txt:
            aba, ref = txt.rsplit("!", 1)
            return aba.strip("'"), ref
        return self.aba, txt

    def _args_brutos(self):
        """Lê argumentos até ')' — intervalos ficam como ('rng', aba, ref)."""
        args = []
        if self._ver() == ("op", ")"):
            self._pegar()
            return args
        while True:
            tipo, txt = self._ver()
            prox = self.toks[self.i + 1] if self.i + 1 < len(self.toks) else (None, None)
            if tipo == "ref" and ":" in txt and prox in (("op", ","), ("op", ")")):
                self._pegar()
                args.append(("rng",) + self._separar(txt))
            else:
                args.append(self.concat())
            if self._pegar()[1] == ")":
                return args

    def funcao(self, nome):
        a = self._args_brutos()
        c = self.calc
        if nome == "SUM":
            total = 0.0
            for x in a:
                vals = c.intervalo(x[1], x[2]) if isinstance(x, tuple) else [x]
                total += sum(v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool))
            return total
        if nome == "LEN":
            return float(len(_txt(a[0])))
        if nome == "RIGHT":
            n = int(_num(a[1])) if len(a) > 1 else 1
            s = _txt(a[0])
            return s[-n:] if n > 0 else ""
        if nome == "SEARCH":
            p = _txt(a[1]).lower().find(_txt(a[0]).lower())
            if p < 0:
                raise _Erro("SEARCH não achou")
            return float(p + 1)
        if nome == "XLOOKUP":
            chave = _txt(a[0]).lower()
            busca, ret = a[1], a[2]
            col_b = c.coluna(busca[1], busca[2].split(":")[0].replace("$", ""))
            col_r = c.coluna(ret[1], ret[2].split(":")[0].replace("$", ""))
            for k, v in zip(col_b, col_r):
                if k is not None and _txt(k).lower() == chave:
                    return v
            return a[3] if len(a) > 3 else None
        raise _Erro(f"função {nome} não suportada")


def gravar_valores(caminho):
    """Calcula as fórmulas do arquivo e grava os valores no próprio .xlsx."""
    wb = openpyxl.load_workbook(caminho)
    calc = Calculadora(wb)
    resultados = []  # um dict ref -> valor por aba, na ordem das abas
    falhas = 0
    for ws in wb.worksheets:
        vals = {}
        for row in ws.iter_rows():
            for cel in row:
                if isinstance(cel.value, str) and cel.value.startswith("="):
                    try:
                        vals[cel.coordinate] = calc.valor(ws.title, cel.coordinate)
                    except (_Erro, KeyError, ValueError, IndexError, TypeError):
                        falhas += 1
        resultados.append(vals)

    # openpyxl grava as abas como xl/worksheets/sheet1.xml, sheet2.xml... na ordem
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx").name
    with zipfile.ZipFile(caminho) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if m and int(m.group(1)) <= len(resultados):
                dados = _injetar(dados.decode("utf-8"), resultados[int(m.group(1)) - 1]).encode("utf-8")
            zout.writestr(item, dados)
    shutil.move(tmp, caminho)
    return falhas


_CEL_FORMULA = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', re.S)


def _injetar(xml, vals):
    def troca(m):
        ref, attrs, f = m.groups()
        if ref not in vals or vals[ref] is None:
            return m.group(0)
        v = vals[ref]
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
        if isinstance(v, str):
            return f'<c r="{ref}"{attrs} t="str"><f>{f}</f><v>{escape(v)}</v></c>'
        if isinstance(v, bool):
            return f'<c r="{ref}"{attrs} t="b"><f>{f}</f><v>{int(v)}</v></c>'
        return f'<c r="{ref}"{attrs}><f>{f}</f><v>{repr(float(v))}</v></c>'
    return _CEL_FORMULA.sub(troca, xml)
